import io
import zipfile
from datetime import datetime
from pathlib import Path

import pandas as pd
import streamlit as st
import streamlit.components.v1 as components
from openpyxl import load_workbook

import core

st.set_page_config(page_title="نظام تسجيل الحضور", page_icon="✅", layout="wide")

scanner = components.declare_component(
    "barcode_scanner", path=str(Path(__file__).parent / "scanner_component")
)

st.markdown(
    """
<style>
html, body, [class*="css"], .stApp { direction: rtl; text-align: right; }
.stTextInput input, .stSelectbox div { text-align: right; }
.result-ok   { background:#e6f4ea; border-right:6px solid #1e8e3e; padding:14px; border-radius:8px; font-size:1.2rem; }
.result-dup  { background:#fef7e0; border-right:6px solid #f9ab00; padding:14px; border-radius:8px; font-size:1.2rem; }
.result-bad  { background:#fce8e6; border-right:6px solid #d93025; padding:14px; border-radius:8px; font-size:1.2rem; }
</style>
""",
    unsafe_allow_html=True,
)

st.title("✅ نظام تسجيل حضور الطلبة بالباركود")

ss = st.session_state
ss.setdefault("book", None)
ss.setdefault("last", None)
ss.setdefault("seen", set())

# ------------------------------------------------------------------ الشريط الجانبي
with st.sidebar:
    st.header("إعداد المحاضرة")
    today = datetime.now(core.TZ).strftime("%Y-%m-%d")

    saved = core.list_saved()
    if saved:
        st.subheader("متابعة ملف محفوظ")
        opts = {f"{m['source_name']} — {m['sheet']}": stem for stem, m in saved}
        pick = st.selectbox("الملف", list(opts.keys()))
        lab2 = st.text_input("اسم عمود المحاضرة (التاريخ)", value=today, key="lab_saved")
        if st.button("فتح وتسجيل المحاضرة", use_container_width=True):
            ss.book = core.open_saved(opts[pick], lab2)
            ss.last, ss.seen = None, set()
            core.autosave(ss.book)
            st.rerun()
        st.divider()

    st.subheader("رفع ملف Excel جديد")
    up = st.file_uploader("ملف الطلبة (xlsx)", type=["xlsx"])
    if up is not None:
        data = up.getvalue()
        sheets = core.sheet_names(data)
        sheet = st.selectbox("الشيت", sheets) if len(sheets) > 1 else sheets[0]
        ws = load_workbook(io.BytesIO(data), read_only=False)[sheet]
        det = core.detect_header(ws)
        hr0, nm0, cd0 = det if det else (1, None, None)
        if det:
            st.success("تم التعرف على صف العناوين والأعمدة تلقائياً")
        else:
            st.warning("لم أتعرف على الأعمدة تلقائياً، حدّدها يدوياً")
        hr = st.number_input("رقم صف العناوين", min_value=1, value=int(hr0), step=1)
        cells = core.header_cells(ws, int(hr))
        if len(cells) >= 2:
            labels = [f"{t}  (عمود {c})" for c, t in cells]
            cols = [c for c, _ in cells]
            ni = cols.index(nm0) if nm0 in cols else 0
            ci = cols.index(cd0) if cd0 in cols else min(1, len(cols) - 1)
            name_col = cols[st.selectbox("عمود اسم الطالب", range(len(cols)), index=ni, format_func=lambda i: labels[i])]
            code_col = cols[st.selectbox("عمود الباركود", range(len(cols)), index=ci, format_func=lambda i: labels[i])]
            label = st.text_input("اسم عمود المحاضرة (التاريخ)", value=today)
            if st.button("ابدأ جلسة الحضور", type="primary", use_container_width=True):
                ss.book = core.AttendanceBook(data, sheet, hr, name_col, code_col, label, source_name=up.name)
                ss.last, ss.seen = None, set()
                core.autosave(ss.book)
                st.rerun()
        else:
            st.error("هذا الصف لا يحتوي عناوين كافية")

B = ss.book

tab_scan, tab_table, tab_cards, tab_help = st.tabs(
    ["📷 المسح", "📋 الجدول والتصدير", "🪪 توليد بطاقات الباركود", "ℹ️ مساعدة"]
)

# ------------------------------------------------------------------ المسح
with tab_scan:
    if B is None:
        st.info("ارفع ملف Excel (أو افتح ملفاً محفوظاً) من الشريط الجانبي.")
    else:
        top = st.container()
        st.caption(f"الملف: **{B.source_name}** — الشيت: **{B.sheet}** — عمود المحاضرة: **{B.label}**")

        mode = st.radio("طريقة المسح", ["كاميرا الهاتف (مستمر)", "ماسح USB / إدخال يدوي"], horizontal=True)
        changed = False

        def apply(code, when=None):
            status, name, t = B.scan(code, when)
            ss.last = (status, name, t, code)

        if mode.startswith("كاميرا"):
            payload = scanner(names=B.names_map(), done=B.done_codes(), key="scanner", default=None)
            if payload:
                for s in payload.get("scans", []):
                    if s["id"] in ss.seen:
                        continue
                    ss.seen.add(s["id"])
                    when = datetime.fromtimestamp(s["ts"] / 1000, core.TZ)
                    if when.timestamp() < B.started - 2:  # مسحات قديمة من جلسة سابقة
                        continue
                    apply(s["code"], when)
                    changed = True
        else:
            def on_enter():
                v = ss.get("manual_code", "")
                if v.strip():
                    apply(v)
                    core.autosave(B)
                ss.manual_code = ""

            st.text_input("امسح الباركود بالماسح أو اكتبه ثم Enter", key="manual_code", on_change=on_enter)

        if changed:
            core.autosave(B)

        present, absent, total = B.counts()
        with top:
            c1, c2, c3 = st.columns(3)
            c1.metric("حاضر", present)
            c2.metric("لم يُسجَّل", absent)
            c3.metric("الإجمالي", total)
            if B.duplicates:
                st.warning(f"يوجد {len(B.duplicates)} باركود مكرر في الملف؛ يُعتمد أول طالب فقط لكل باركود.")
            if ss.last:
                status, name, t, code = ss.last
                if status == "ok":
                    st.markdown(f'<div class="result-ok">✅ تم تسجيل حضور: <b>{name}</b><br>الوقت: {t}</div>', unsafe_allow_html=True)
                elif status == "dup":
                    st.markdown(f'<div class="result-dup">⚠️ <b>{name}</b> مسجَّل مسبقاً في: {t}</div>', unsafe_allow_html=True)
                else:
                    st.markdown(f'<div class="result-bad">❌ الباركود <b>{code}</b> غير موجود في قائمة هذا الشيت</div>', unsafe_allow_html=True)

# ------------------------------------------------------------------ الجدول
with tab_table:
    if B is None:
        st.info("لا توجد جلسة بعد.")
    else:
        present, absent, total = B.counts()
        st.write(f"الحضور: **{present}** من **{total}**")
        rows = B.table()
        show = pd.DataFrame(
            {
                "حاضر؟": [bool(r["time"]) for r in rows],
                "الاسم": [r["name"] for r in rows],
                "الباركود": [r["code"] for r in rows],
                f"الوقت ({B.label})": [r["time"] for r in rows],
            }
        )
        edited = st.data_editor(
            show, use_container_width=True, hide_index=True,
            disabled=["الاسم", "الباركود", f"الوقت ({B.label})"], key="editor",
        )
        if st.button("تطبيق التعديل اليدوي على الحضور"):
            for i, r in enumerate(rows):
                now_p = bool(edited.iloc[i]["حاضر؟"])
                if now_p != bool(r["time"]):
                    B.set_manual(r["row"], now_p)
            core.autosave(B)
            st.rerun()

        stem = Path(B.source_name).stem
        st.download_button(
            "⬇️ تنزيل ملف Excel المحدَّث (نفس الشيت)",
            data=B.to_bytes(),
            file_name=f"{stem}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            type="primary",
        )
        st.download_button("⬇️ تنزيل سجل المسح (CSV)", data=B.log_csv(), file_name=f"سجل_{stem}.csv", mime="text/csv")
        st.caption("الملف المحدَّث هو نفس ملفك الأصلي مع عمود جديد لهذه المحاضرة، ويحتفظ بالتنسيق وبكل المحاضرات السابقة.")

# ------------------------------------------------------------------ البطاقات
with tab_cards:
    st.write("إذا لم يكن لدى الطلبة باركود جاهز، ارفع ملفاً يحتوي الاسم والرقم (مثل الرقم الجامعي) لتوليد صورة باركود لكل طالب.")
    cu = st.file_uploader("ملف الطلبة", type=["xlsx"], key="cards_up")
    kind = st.radio("نوع الرمز", ["QR", "Code128"], horizontal=True)
    if cu is not None:
        cdata = cu.getvalue()
        cws = load_workbook(io.BytesIO(cdata))[core.sheet_names(cdata)[0]]
        det = core.detect_header(cws)
        chr_ = det[0] if det else 1
        ccells = core.header_cells(cws, chr_)
        if len(ccells) >= 2:
            lab = [f"{t} (عمود {c})" for c, t in ccells]
            n_i = st.selectbox("عمود الاسم", range(len(ccells)), format_func=lambda i: lab[i], key="cn")
            v_i = st.selectbox("عمود القيمة التي يحملها الباركود", range(len(ccells)), index=1, format_func=lambda i: lab[i], key="cv")
            if st.button("توليد الصور (ملف ZIP)"):
                buf = io.BytesIO()
                with zipfile.ZipFile(buf, "w") as z:
                    for r in range(chr_ + 1, cws.max_row + 1):
                        val = core.norm(cws.cell(r, ccells[v_i][0]).value)
                        nm = str(cws.cell(r, ccells[n_i][0]).value or "").strip()
                        if val:
                            safe = "".join(ch for ch in nm if ch not in '\\/:*?"<>|').strip() or val
                            z.writestr(f"{safe}_{val}.png", core.make_code_image(val, kind))
                st.download_button("⬇️ تنزيل ZIP", buf.getvalue(), "barcodes.zip", "application/zip")

# ------------------------------------------------------------------ مساعدة
with tab_help:
    st.markdown(
        """
**خطوات الاستخدام**
1. ارفع ملف Excel الخاص بالشعبة كما هو من النظام (نفس الشيت). يتعرف التطبيق على صف العناوين وعمودي الاسم والباركود، ويمكنك تصحيحهما.
2. اكتب اسم عمود المحاضرة (التاريخ) واضغط **ابدأ جلسة الحضور**.
3. في تبويب **المسح** اضغط «تشغيل الكاميرا» ثم مرّر باركود كل طالب أمام الكاميرا؛ تُقرأ الأكواد تباعاً دون التقاط صور، ويُكتب وقت الحضور مباشرة في خانة الطالب.
4. يُحفظ كل مسح تلقائياً على الخادم، وفي المحاضرة التالية افتح الملف من «متابعة ملف محفوظ» فيُضاف عمود جديد بجانب السابق.
5. نزّل ملف Excel المحدَّث من تبويب **الجدول والتصدير** وارفعه لنظام الجامعة أو احتفظ به.

**ملاحظات**
- التسجيل يتم داخل ملفك الأصلي نفسه، فلا يتغير الشيت ولا التنسيق، ويُضاف عمود جديد فقط لكل محاضرة.
- الملفات المحفوظة على الخادم قد تُمسح عند إعادة تشغيل الاستضافة المجانية، فنزّل نسخة من الملف بين وقت وآخر.
- الكاميرا تحتاج رابط https وإذن المتصفح.
- ماسح USB يعمل كلوحة مفاتيح: اختر «ماسح USB» وامسح مباشرة.
- صيغة xls القديمة غير مدعومة؛ احفظ الملف بصيغة xlsx أولاً.
"""
    )
