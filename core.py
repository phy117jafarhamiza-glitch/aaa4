"""منطق تطبيق تسجيل الحضور.

يعمل مباشرة داخل ملف Excel الأصلي (نفس الشيت ونفس التنسيق) عبر openpyxl:
يُضاف عمود جديد لكل محاضرة في نهاية صف العناوين، ويُكتب وقت المسح أمام كل طالب.
"""
import io
import json
import os
import re
from copy import copy
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

TZ = ZoneInfo("Asia/Baghdad")

NAME_HINTS = ["اسم الطالب", "الاسم", "name", "student"]
CODE_HINTS = ["الباركود", "باركود", "barcode", "qr", "code", "الرقم الجامعي", "رقم الطالب", "id"]

DATA_DIR = Path(os.environ.get("ATTENDANCE_DATA_DIR", "attendance_data"))


# ------------------------------------------------------------------ أدوات
def norm(value) -> str:
    """توحيد شكل الباركود: نص بلا مسافات، وإزالة .0 التي يضيفها Excel للأرقام."""
    if value is None:
        return ""
    s = str(value).strip()
    if re.fullmatch(r"\d+\.0+", s):
        s = s.split(".")[0]
    return re.sub(r"\s+", "", s).upper()


def _txt(v) -> str:
    return "" if v is None else str(v).strip()


def sheet_names(data: bytes) -> list:
    wb = load_workbook(io.BytesIO(data), read_only=True)
    names = list(wb.sheetnames)
    wb.close()
    return names


def _find(vals, hints, exclude=()):
    """vals: قائمة (رقم العمود، النص). يعيد رقم العمود الأنسب لأحد المفاتيح."""
    for h in hints:
        for col, t in vals:
            if col not in exclude and t.lower() == h.lower():
                return col
    for h in hints:
        for col, t in vals:
            if col not in exclude and h.lower() in t.lower():
                return col
    return None


def detect_header(ws, max_rows: int = 40):
    """يبحث عن صف العناوين (قد تسبقه صفوف عنوان الجامعة/الكلية).
    يعيد (صف العناوين، عمود الاسم، عمود الباركود) أو None."""
    for r in range(1, min(ws.max_row, max_rows) + 1):
        vals = [(c.column, _txt(c.value)) for c in ws[r] if _txt(c.value)]
        if len(vals) < 2:
            continue
        cd = _find(vals, CODE_HINTS)
        nm = _find(vals, NAME_HINTS, exclude=(cd,))
        if nm and cd:
            return r, nm, cd
    return None


def header_cells(ws, header_row: int) -> list:
    return [(c.column, _txt(c.value)) for c in ws[header_row] if _txt(c.value)]


def make_code_image(value: str, kind: str = "QR") -> bytes:
    """يولّد صورة PNG لباركود/QR لقيمة معينة."""
    buf = io.BytesIO()
    if kind == "QR":
        import qrcode

        qrcode.make(value, box_size=8, border=2).save(buf, format="PNG")
    else:
        import barcode
        from barcode.writer import ImageWriter

        barcode.get("code128", value, writer=ImageWriter()).write(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ الدفتر
class AttendanceBook:
    """ملف Excel واحد (شيت واحد) مع عمود حضور للمحاضرة الحالية."""

    def __init__(self, data, sheet, header_row, name_col, code_col, label,
                 source_name="attendance.xlsx", log=None):
        self.source_name = source_name
        self.sheet = sheet
        self.header_row = int(header_row)
        self.name_col = int(name_col)
        self.code_col = int(code_col)
        self.label = _txt(label) or datetime.now(TZ).strftime("%Y-%m-%d")
        self.log = list(log or [])  # (وقت كامل، باركود، نتيجة، اسم)
        self.started = datetime.now(TZ).timestamp()

        self.wb = load_workbook(io.BytesIO(data))
        self.ws = self.wb[sheet]

        self.rows = {}        # باركود موحّد -> رقم الصف
        self.students = []    # (صف، اسم، باركود أصلي)
        self.duplicates = set()
        for r in range(self.header_row + 1, self.ws.max_row + 1):
            code = norm(self.ws.cell(r, self.code_col).value)
            name = _txt(self.ws.cell(r, self.name_col).value)
            if not code and not name:
                continue
            self.students.append((r, name, _txt(self.ws.cell(r, self.code_col).value)))
            if not code:
                continue
            if code in self.rows:
                self.duplicates.add(code)
            else:
                self.rows[code] = r

        self.col = self._find_or_create_col()

    # -------- عمود المحاضرة
    def _find_or_create_col(self) -> int:
        ws, hr = self.ws, self.header_row
        last = 0
        for c in ws[hr]:
            if _txt(c.value):
                last = max(last, c.column)
                if _txt(c.value) == self.label:
                    return c.column
        new = last + 1
        head = ws.cell(hr, new)
        head.value = self.label
        src = ws.cell(hr, last)
        head._style = copy(src._style)
        for r, _, _ in self.students:
            ws.cell(r, new)._style = copy(ws.cell(r, last)._style)
        ws.column_dimensions[head.column_letter].width = max(12, len(self.label) + 4)
        return new

    # -------- المسح
    def scan(self, raw_code, when=None):
        """يعيد (الحالة, الاسم, الوقت): ok | dup | unknown."""
        code = norm(raw_code)
        when = when or datetime.now(TZ)
        full = when.strftime("%Y-%m-%d %H:%M:%S")
        if not code:
            return "unknown", "", ""
        r = self.rows.get(code)
        if r is None:
            self.log.append((full, code, "غير موجود", ""))
            return "unknown", "", full
        name = _txt(self.ws.cell(r, self.name_col).value)
        cell = self.ws.cell(r, self.col)
        if _txt(cell.value):
            self.log.append((full, code, "مكرر", name))
            return "dup", name, _txt(cell.value)
        cell.value = when.strftime("%H:%M:%S")
        self.log.append((full, code, "تم التسجيل", name))
        return "ok", name, cell.value

    def set_manual(self, row: int, present: bool):
        cell = self.ws.cell(row, self.col)
        if present:
            cell.value = datetime.now(TZ).strftime("%H:%M:%S")
            self.log.append((datetime.now(TZ).strftime("%Y-%m-%d %H:%M:%S"), "", "يدوي",
                             _txt(self.ws.cell(row, self.name_col).value)))
        else:
            cell.value = None

    # -------- عرض
    def table(self) -> list:
        """صفوف للعرض: dict(row, name, code, time)."""
        return [
            {"row": r, "name": n, "code": c, "time": _txt(self.ws.cell(r, self.col).value)}
            for r, n, c in self.students
        ]

    def counts(self):
        present = sum(1 for r, _, _ in self.students if _txt(self.ws.cell(r, self.col).value))
        return present, len(self.students) - present, len(self.students)

    def names_map(self) -> dict:
        return {k: _txt(self.ws.cell(r, self.name_col).value) for k, r in self.rows.items()}

    def done_codes(self) -> list:
        return [k for k, r in self.rows.items() if _txt(self.ws.cell(r, self.col).value)]

    # -------- حفظ
    def to_bytes(self) -> bytes:
        buf = io.BytesIO()
        self.wb.save(buf)
        return buf.getvalue()

    def log_csv(self) -> bytes:
        import csv

        out = io.StringIO()
        w = csv.writer(out)
        w.writerow(["وقت المسح", "الباركود", "النتيجة", "اسم الطالب"])
        w.writerows(self.log)
        return ("﻿" + out.getvalue()).encode("utf-8")

    def meta(self) -> dict:
        return {
            "sheet": self.sheet, "header_row": self.header_row,
            "name_col": self.name_col, "code_col": self.code_col,
            "label": self.label, "source_name": self.source_name,
            "log": self.log,
        }


# ------------------------------------------------------------------ الحفظ التلقائي
def _safe(s: str) -> str:
    return re.sub(r"[^\w\-]+", "_", s, flags=re.UNICODE).strip("_") or "file"


def _stem(book: AttendanceBook) -> str:
    return f"{_safe(Path(book.source_name).stem)}__{_safe(book.sheet)}"


def autosave(book: AttendanceBook):
    """يحفظ الملف المحدّث وبيانات الجلسة على القرص بعد كل مسح."""
    try:
        DATA_DIR.mkdir(parents=True, exist_ok=True)
        stem = _stem(book)
        (DATA_DIR / f"{stem}.xlsx").write_bytes(book.to_bytes())
        (DATA_DIR / f"{stem}.json").write_text(
            json.dumps(book.meta(), ensure_ascii=False), encoding="utf-8"
        )
        return True
    except Exception:
        return False


def list_saved() -> list:
    if not DATA_DIR.exists():
        return []
    out = []
    for j in sorted(DATA_DIR.glob("*.json"), key=lambda p: p.stat().st_mtime, reverse=True):
        if j.with_suffix(".xlsx").exists():
            try:
                m = json.loads(j.read_text(encoding="utf-8"))
                out.append((j.stem, m))
            except Exception:
                pass
    return out


def open_saved(stem: str, label: str) -> AttendanceBook:
    m = json.loads((DATA_DIR / f"{stem}.json").read_text(encoding="utf-8"))
    data = (DATA_DIR / f"{stem}.xlsx").read_bytes()
    return AttendanceBook(
        data, m["sheet"], m["header_row"], m["name_col"], m["code_col"], label,
        source_name=m["source_name"], log=[tuple(x) for x in m.get("log", [])],
    )
